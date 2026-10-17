from pathlib import Path

from docx import Document as WordDocument
from docx.shared import Inches

from ocr_to_docs.models import Block, BlockType, Document


def export_docx(document: Document, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    word = WordDocument()

    for page_index, page in enumerate(document.pages):
        if page_index:
            word.add_page_break()
        for block in page.blocks:
            _add_block(word, block)

    word.save(output)


def _add_block(word: WordDocument, block: Block) -> None:
    if block.kind == BlockType.HEADING:
        word.add_heading(block.text, level=1)
    elif block.kind == BlockType.TABLE and (rows := _markdown_table(block.text)):
        table = word.add_table(rows=len(rows), cols=max(map(len, rows)))
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    elif block.kind in {BlockType.CHART, BlockType.IMAGE}:
        if block.image_path:
            word.add_picture(str(block.image_path), width=Inches(6))
        if block.text:
            word.add_paragraph(block.text)
    elif block.kind == BlockType.FORMULA:
        word.add_paragraph(block.text, style="Quote")
    elif block.text:
        word.add_paragraph(block.text)


def _markdown_table(text: str) -> list[list[str]]:
    rows = [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in text.splitlines()
        if "|" in line
    ]
    if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
        rows.pop(1)
    return rows
