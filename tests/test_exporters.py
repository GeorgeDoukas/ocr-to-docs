from pathlib import Path

import pymupdf
from docx import Document as WordDocument
from PIL import Image

from ocr_to_docs.exporters import export_docx, export_markdown, export_searchable_pdf
from ocr_to_docs.models import Block, BlockType, BoundingBox, Document, Page


def sample_document(tmp_path: Path) -> Document:
    page_image = tmp_path / "page.png"
    chart_image = tmp_path / "chart.png"
    Image.new("RGB", (200, 100), "white").save(page_image)
    Image.new("RGB", (80, 40), "blue").save(chart_image)
    return Document(
        source=str(page_image),
        pages=[
            Page(
                number=1,
                width=200,
                height=100,
                source_image=page_image,
                blocks=[
                    Block(BlockType.HEADING, "Report", BoundingBox(10, 5, 150, 20)),
                    Block(BlockType.TEXT, "English and Ελληνικά", BoundingBox(10, 25, 190, 45)),
                    Block(
                        BlockType.TABLE,
                        "| Year | Sales |\n|---|---|\n| 2026 | 10 |",
                        BoundingBox(10, 50, 190, 75),
                    ),
                    Block(
                        BlockType.CHART,
                        "2026: 10",
                        BoundingBox(10, 76, 90, 98),
                        image_path=chart_image,
                    ),
                ],
            )
        ],
    )


def test_markdown_export_preserves_structure_and_chart(tmp_path: Path) -> None:
    output = tmp_path / "result.md"

    export_markdown(sample_document(tmp_path), output)

    text = output.read_text(encoding="utf-8")
    assert "# Report" in text
    assert "| Year | Sales |" in text
    assert "![Chart]" in text
    assert (tmp_path / "result_assets" / "chart.png").exists()


def test_docx_export_creates_native_table(tmp_path: Path) -> None:
    output = tmp_path / "result.docx"

    export_docx(sample_document(tmp_path), output)

    word = WordDocument(output)
    assert "Report" in [paragraph.text for paragraph in word.paragraphs]
    assert word.tables[0].cell(1, 1).text == "10"


def test_searchable_pdf_contains_invisible_ocr_text(tmp_path: Path) -> None:
    output = tmp_path / "result.pdf"

    export_searchable_pdf(sample_document(tmp_path), output)

    with pymupdf.open(output) as pdf:
        text = pdf[0].get_text()
    assert "English and Ελληνικά" in text


def test_searchable_pdf_respects_image_dpi(tmp_path: Path) -> None:
    source = tmp_path / "letter.png"
    Image.new("RGB", (300, 600), "white").save(source, dpi=(300, 300))
    document = Document(
        source=str(source),
        pages=[
            Page(
                number=1,
                width=300,
                height=600,
                physical_width=72,
                physical_height=144,
                source_image=source,
            )
        ],
    )
    output = tmp_path / "letter.pdf"

    export_searchable_pdf(document, output)

    with pymupdf.open(output) as pdf:
        assert pdf[0].rect.width == 72
        assert pdf[0].rect.height == 144
